#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""生成素材拆解分析库的标准 Excel 模板（多游戏版 + 真实投放数据列）"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()

# ── 样式 ──
header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="2E4057", end_color="2E4057", fill_type="solid")
tag_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
data_font = Font(name="微软雅黑", size=10)
thin_border = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)

def style_header(ws, row=1):
    for col in range(1, ws.max_column + 1):
        c = ws.cell(row=row, column=col)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = thin_border

def style_data(ws, start_row=2):
    for row in range(start_row, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            c = ws.cell(row=row, column=col)
            c.font = data_font
            c.border = thin_border
            c.alignment = Alignment(vertical="top", wrap_text=True)

def auto_width(ws, max_w=40):
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        ml = max((len(str(ws.cell(row=r, column=col).value or ""))
                  for r in range(1, min(ws.max_row + 1, 20))), default=8)
        ws.column_dimensions[letter].width = min(max(ml + 2, 10), max_w)


# =============================================================
# Sheet 1: 素材总览（输入表 — 对齐真实投放数据列）
# =============================================================
ws1 = wb.active
ws1.title = "素材总览"

ws1.append([
    "素材名称",         # A — Ad（清理后缀后）
    "游戏名称",         # B — 多游戏核心字段
    "平台",             # C
    "类型",             # D — 视频/图片
    "Campaign",         # E
    "链接",             # F
    "入库时间",         # G
    "状态",             # H — 待分析/AI已分析/人工确认/已上线/已淘汰
    "上传人",           # I
    # ── 投放数据（来自 1.xlsx，列名对齐） ──
    "花费",             # J — 花费（美元）
    "点击率",           # K — CTR
    "安装",             # L — 安装数
    "转化率",           # M — 转化率
    "单价",             # N — CPI（美元）
    "首充用户",         # O — 首充用户数
    "首充率",           # P — 首充率
    "首日回收",         # Q — 首日回收（美元）
    "单个付费用户成本", # R — 单个付费用户成本
    "首日ROI",          # S — D1 ROI
    "累计付费",         # T — 累计付费（美元）
    "累计ROI",          # U — 累计 ROI
    "次留",             # V — 次日留存率
    # ── 补充 ──
    "时长(秒)",         # W
    "分辨率",           # X
    "备注",             # Y
])

style_header(ws1)
style_data(ws1)
auto_width(ws1)
ws1.freeze_panes = "A2"


# =============================================================
# Sheet 2: 逐帧拆解（输出表）
# =============================================================
ws2 = wb.create_sheet("逐帧拆解")

ws2.append([
    "素材名称", "游戏名称", "时间轴", "画面内容",
    "画面元素", "文案/字幕", "音效", "目的", "可借鉴点",
])

style_header(ws2)
style_data(ws2)
auto_width(ws2)
ws2.freeze_panes = "A2"


# =============================================================
# Sheet 3: 总结分析（输出表）
# =============================================================
ws3 = wb.create_sheet("总结分析")

ws3.append([
    "素材名称", "游戏名称", "素材类型",
    "Hook内容", "Hook类型", "人物", "UI元素", "特效",
    "色彩", "字幕", "节奏", "BGM", "核心卖点",
    "hook_策略", "角色元素", "场景元素", "制作方式", "叙事模式", "画面风格",
    "目标受众", "情绪触发", "营销策略", "可优化方向",
    "标签状态", "AI置信度", "分析时间", "人工批注",
])

style_header(ws3)
style_data(ws3)
for col in range(14, 20):
    for row in range(2, 40):
        ws3.cell(row=row, column=col).fill = tag_fill
auto_width(ws3)
ws3.freeze_panes = "D2"


# =============================================================
# Sheet 4: 标签字典（多游戏）
# =============================================================
ws4 = wb.create_sheet("标签字典")

ws4.append(["游戏名称", "维度", "标签值", "说明", "示例素材"])

common = [
    ("通用", "hook_策略", "反常识标题", "用反常识/信息差制造悬念", "千万别选反弹球"),
    ("通用", "hook_策略", "教学攻略", "教用户怎么玩/阵容推荐", "阵容搭配教学"),
    ("通用", "hook_策略", "福利展示", "突出免费福利/抽卡资源", "登录就送XX抽"),
    ("通用", "hook_策略", "擦边/角色吸引", "用角色立绘/身材筛选用户", "猫耳立绘特写"),
    ("通用", "hook_策略", "剧情/抓马", "用短剧情吸引追看", "逆袭打脸"),
    ("通用", "hook_策略", "口播推荐", "真人口播/AI口播推荐", "今天给大家推荐"),
    ("通用", "hook_策略", "问题型标题", "抛出问题引发好奇", "你还在用XX队？"),
    ("通用", "hook_策略", "选择互动", "通过选择引导剧情", "你会选哪个？"),
    ("通用", "hook_策略", "福利激励", "转盘抽奖/签到福利", "转盘抽奖"),
    ("通用", "hook_策略", "核心玩法展示", "直接展示核心玩法", "实机战斗画面"),
    ("通用", "hook_策略", "来电钩子", "以来电/消息通知开场", "来电界面"),
    ("通用", "hook_策略", "定制互动", "强调可定制外观/剧情", "定制你的角色"),
    ("通用", "hook_策略", "噱头/爆点", "以噱头内容吸引注意力", "爆点标题"),
    ("通用", "角色元素", "黑皮", "深色皮肤角色展示", ""),
    ("通用", "角色元素", "猫耳/兽耳", "猫耳/兽耳角色展示", ""),
    ("通用", "角色元素", "Q版角色", "Q版像素/Q版小人", ""),
    ("通用", "角色元素", "多角色", "多个角色同时展示", ""),
    ("通用", "角色元素", "立绘展示", "角色大立绘/特写", ""),
    ("通用", "角色元素", "真人出镜", "真人出镜拍摄", ""),
    ("通用", "角色元素", "AI虚拟人", "AI生成虚拟人物", ""),
    ("通用", "角色元素", "无角色", "纯游戏画面", ""),
    ("通用", "场景元素", "弹球/战斗", "核心战斗画面", ""),
    ("通用", "场景元素", "选择界面", "角色选择/阵容搭配", ""),
    ("通用", "场景元素", "Boss战", "Boss出场/战斗", ""),
    ("通用", "场景元素", "胜利结算", "通关/胜利/奖励", ""),
    ("通用", "场景元素", "福利界面", "抽卡/签到/奖励", ""),
    ("通用", "场景元素", "模拟UI", "模拟聊天/约会软件", ""),
    ("通用", "场景元素", "直播间", "模拟直播场景", ""),
    ("通用", "制作方式", "K帧动画", "关键帧动画", ""),
    ("通用", "制作方式", "真人口播", "真人出镜口播", ""),
    ("通用", "制作方式", "AI口播", "AI虚拟人/语音合成", ""),
    ("通用", "制作方式", "实机录屏", "真实游戏录屏", ""),
    ("通用", "制作方式", "图片轮播", "静态图片+转场", ""),
    ("通用", "制作方式", "混合", "多种方式混合", ""),
    ("通用", "叙事模式", "问题-解法", "抛出问题→展示解法", ""),
    ("通用", "叙事模式", "反差逆袭", "弱者→变强打脸", ""),
    ("通用", "叙事模式", "攻略教学", "Step-by-step教玩法", ""),
    ("通用", "叙事模式", "角色推荐", "推荐角色→展示强度", ""),
    ("通用", "叙事模式", "悬念-揭秘", "制造悬念→爽点释放", ""),
    ("通用", "叙事模式", "福利直给", "直接展示福利", ""),
    ("通用", "叙事模式", "选项驱动", "选择引导剧情", ""),
    ("通用", "叙事模式", "擦边吸引", "以擦边内容吸引", ""),
    ("通用", "叙事模式", "角色展示", "角色立绘/自我介绍", ""),
    ("通用", "叙事模式", "聊天模拟", "模拟聊天/约会软件", ""),
    ("通用", "画面风格", "二次元精致渲染", "日系二次元高细节", ""),
    ("通用", "画面风格", "Q版可爱风", "Q版/像素风", ""),
    ("通用", "画面风格", "AI合成", "AI生成/合成", ""),
    ("通用", "画面风格", "实机录屏", "直接录屏/截图", ""),
    ("通用", "画面风格", "实拍", "真人实拍", ""),
    ("通用", "画面风格", "混剪", "多种风格混合", ""),
]

for tag in common:
    ws4.append(list(tag))

style_header(ws4)
style_data(ws4)
auto_width(ws4)
ws4.freeze_panes = "A2"


# =============================================================
# Sheet 5: 效果聚合（输出表）
# =============================================================
ws5 = wb.create_sheet("效果聚合")

ws5.append([
    "游戏名称", "分析维度", "标签值", "素材数",
    "平均花费", "平均单价", "平均首日ROI", "平均累计ROI",
    "平均次留", "结论",
])

style_header(ws5)
style_data(ws5)
auto_width(ws5)
ws5.freeze_panes = "A2"


# =============================================================
# 保存
# =============================================================
ROOT = Path(__file__).resolve().parent
output = str(ROOT / "素材拆解分析模板.xlsx")
wb.save(output)
print(f"✅ 模板已生成: {output}")
print(f"   素材总览列已对齐 1.xlsx 的真实投放数据格式")
