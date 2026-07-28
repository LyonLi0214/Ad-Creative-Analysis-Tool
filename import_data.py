#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
import_data.py — 从投放数据表(1.xlsx)导入素材到素材总览

用法：
    python import_data.py                      # 导入全部
    python import_data.py --game 心动防线       # 指定游戏名称
    python import_data.py --source C:\path\to\1.xlsx
"""
import re
import argparse
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook

XLSX_PATH = Path(__file__).parent / "素材拆解分析模板.xlsx"

def clean_ad_name(raw_name):
    """清理 Ad 名称：去掉 _广告名称... 后缀"""
    if not raw_name:
        return ""
    name = str(raw_name).strip()
    # 去掉 _广告名称YYYY-MM-DD HH:MM:SS 后缀
    name = re.sub(r'_广告名称\d{4}-\d{2}-\d{2}.*$', '', name)
    # 去掉末尾的 .mp4
    name = re.sub(r'\.mp4$', '', name)
    return name

def detect_game(name):
    """从素材名推断游戏名称"""
    if '心动' in name:
        return '心动防线'
    if '黑帮' in name or '黑幫' in name:
        return '黑帮'
    return ''

def detect_type(name):
    """从素材名推断类型"""
    if name.startswith('V') or name.endswith('.mp4'):
        return '视频'
    if name.startswith('P'):
        return '图片'
    return '视频'

def detect_platform(name):
    """从素材名推断平台（粗略）"""
    # 大部分是 Google/FB，暂无法精确判断
    return ''

def main():
    parser = argparse.ArgumentParser(description="导入投放数据到素材总览")
    parser.add_argument("--source", default="1.xlsx", help="源数据文件路径")
    parser.add_argument("--game", default="", help="指定游戏名称（不填则自动推断）")
    args = parser.parse_args()

    source_path = Path(args.source)
    if not source_path.exists():
        print(f"❌ 源文件不存在: {source_path}")
        return

    # 读源数据
    src_wb = load_workbook(str(source_path), data_only=True)
    src_ws = src_wb.active

    # 读取表头
    headers = {}
    for col in range(1, src_ws.max_column + 1):
        h = src_ws.cell(row=1, column=col).value
        if h:
            headers[str(h).strip()] = col

    print(f"源文件: {source_path.name}")
    print(f"表头: {list(headers.keys())}")
    print(f"数据行: {src_ws.max_row - 1}")

    # 读取数据
    rows_data = []
    for row in range(2, src_ws.max_row + 1):
        ad_raw = src_ws.cell(row=row, column=headers.get('Ad', 1)).value
        if not ad_raw:
            continue

        ad_name = clean_ad_name(ad_raw)
        game = args.game or detect_game(ad_name)

        data = {
            '素材名称': ad_name,
            '游戏名称': game,
            '类型': detect_type(ad_name),
            '平台': detect_platform(ad_name),
            '状态': '待分析',
            '花费': src_ws.cell(row=row, column=headers.get('花费', 2)).value,
            '点击率': src_ws.cell(row=row, column=headers.get('点击率', 3)).value,
            '安装': src_ws.cell(row=row, column=headers.get('安装', 4)).value,
            '转化率': src_ws.cell(row=row, column=headers.get('转化率', 5)).value,
            '单价': src_ws.cell(row=row, column=headers.get('单价', 6)).value,
            '首充用户': src_ws.cell(row=row, column=headers.get('首充用户', 7)).value,
            '首充率': src_ws.cell(row=row, column=headers.get('首充率', 8)).value,
            '首日回收': src_ws.cell(row=row, column=headers.get('首日回收', 9)).value,
            '单个付费用户成本': src_ws.cell(row=row, column=headers.get('单个付费用户成本', 10)).value,
            '首日ROI': src_ws.cell(row=row, column=headers.get('首日ROI', 11)).value,
            '累计付费': src_ws.cell(row=row, column=headers.get('累计付费', 12)).value,
            '累计ROI': src_ws.cell(row=row, column=headers.get('累计ROI', 13)).value,
            '次留': src_ws.cell(row=row, column=headers.get('次留', 14)).value,
        }
        rows_data.append(data)

    src_wb.close()

    print(f"\n解析到 {len(rows_data)} 条素材")

    # 写入模板
    dst_wb = load_workbook(str(XLSX_PATH))
    dst_ws = dst_wb["素材总览"]

    # 获取目标表头
    dst_headers = {}
    for col in range(1, dst_ws.max_column + 1):
        h = dst_ws.cell(row=1, column=col).value
        if h:
            dst_headers[str(h).strip()] = col

    # 检查已存在的素材名（避免重复）
    existing = set()
    for row in range(2, dst_ws.max_row + 1):
        name = dst_ws.cell(row=row, column=dst_headers.get('素材名称', 1)).value
        if name:
            existing.add(str(name).strip())

    # 写入新数据
    added = 0
    skipped = 0
    for data in rows_data:
        if data['素材名称'] in existing:
            skipped += 1
            continue

        row_num = dst_ws.max_row + 1
        for header, col in dst_headers.items():
            dst_ws.cell(row=row_num, column=col).value = data.get(header)
        added += 1
        existing.add(data['素材名称'])

    dst_wb.save(str(XLSX_PATH))
    dst_wb.close()

    print(f"写入素材总览: {added} 条新增, {skipped} 条跳过（已存在）")
    print(f"\n下一步: 运行 python agent_tagger.py 开始分析")


if __name__ == "__main__":
    main()
