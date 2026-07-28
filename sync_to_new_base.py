import requests, json, time, os
from pathlib import Path
from openpyxl import load_workbook

# Load .env
for line in Path(__file__).parent.joinpath('.env').read_text(encoding='utf-8').splitlines():
    line = line.strip()
    if line and not line.startswith('#') and '=' in line:
        k, v = line.split('=', 1)
        os.environ.setdefault(k.strip(), v.strip())

H = {'xc-token': os.environ['NOCODB_TOKEN'], 'Content-Type': 'application/json'}
BASE = os.environ['NOCODB_BASE']
T_ASSETS = os.environ['NOCODB_TABLE_ASSETS']
T_DICT = os.environ['NOCODB_TABLE_DICT']
T_TIMELINE = os.environ['NOCODB_TABLE_TIMELINE']
T_ANALYSIS = os.environ['NOCODB_TABLE_ANALYSIS']

XLSX = Path(__file__).parent / '素材拆解分析模板.xlsx'

def insert_records(table_id, rows):
    url = f'{BASE}/api/v2/tables/{table_id}/records'
    r = requests.post(url, headers=H, json=rows, timeout=30)
    if r.ok:
        return len(rows)
    else:
        print(f'  ! Insert failed: {r.text[:150]}')
        return 0

# ═══════════════════════════════════════════════════════════
# 1. 填充标签字典
# ═══════════════════════════════════════════════════════════
print('=== 填充标签字典 ===')
wb = load_workbook(str(XLSX), data_only=True)
ws_dict = wb['标签字典']

tags = []
for row in range(2, ws_dict.max_row + 1):
    game = ws_dict.cell(row=row, column=1).value
    dim = ws_dict.cell(row=row, column=2).value
    val = ws_dict.cell(row=row, column=3).value
    desc = ws_dict.cell(row=row, column=4).value
    example = ws_dict.cell(row=row, column=5).value
    if val:
        tags.append({
            '游戏名称': game or '',
            '维度': dim or '',
            '标签值': val,
            '说明': desc or '',
            '示例素材': example or '',
        })

ok = 0
for i in range(0, len(tags), 20):
    batch = tags[i:i+20]
    ok += insert_records(T_DICT, batch)
    time.sleep(0.3)
print(f'  标签字典: {ok}/{len(tags)} 条')

# ═══════════════════════════════════════════════════════════
# 2. 同步素材总览
# ═══════════════════════════════════════════════════════════
print('\n=== 同步素材总览 ===')
ws_assets = wb['素材总览']
headers = [ws_assets.cell(row=1, column=c).value for c in range(1, ws_assets.max_column + 1)]
col_map = {h: i + 1 for i, h in enumerate(headers)}

COL_MAP = {
    '素材名称': '素材名称', '游戏名称': '游戏名称', '平台': '平台',
    '类型': '类型', 'Campaign': 'Campaign', '链接': '链接',
    '入库时间': '入库时间', '状态': '状态', '上传人': '上传人',
    '花费': '花费', '点击率': '点击率', '安装': '安装',
    '转化率': '转化率', '单价': '单价', '首充用户': '首充用户',
    '首充率': '首充率', '首日回收': '首日回收',
    '单个付费用户成本': '单个付费用户成本', '首日ROI': '首日ROI',
    '累计付费': '累计付费', '累计ROI': '累计ROI', '次留': '次留',
    '时长(秒)': '时长', '分辨率': '分辨率', '备注': '备注',
}

rows = []
for row in range(2, ws_assets.max_row + 1):
    name = ws_assets.cell(row=row, column=col_map.get('素材名称', 1)).value
    if not name:
        continue
    record = {}
    for excel_col, noco_col in COL_MAP.items():
        ci = col_map.get(excel_col)
        if ci:
            val = ws_assets.cell(row=row, column=ci).value
            if val is not None:
                record[noco_col] = val
    rows.append(record)

ok = 0
for i in range(0, len(rows), 20):
    batch = rows[i:i+20]
    ok += insert_records(T_ASSETS, batch)
    time.sleep(0.3)
print(f'  素材主表: {ok}/{len(rows)} 条')

# ═══════════════════════════════════════════════════════════
# 3. 同步逐帧拆解 + 总结分析
# ═══════════════════════════════════════════════════════════
print('\n=== 同步逐帧拆解 ===')
ws_tl = wb['逐帧拆解']
tl_rows = []
for row in range(2, ws_tl.max_row + 1):
    name = ws_tl.cell(row=row, column=1).value
    if not name:
        continue
    tl_rows.append({
        '素材名称': ws_tl.cell(row=row, column=1).value or '',
        '游戏名称': ws_tl.cell(row=row, column=2).value or '',
        '时间轴': ws_tl.cell(row=row, column=3).value or '',
        '画面内容': ws_tl.cell(row=row, column=4).value or '',
        '画面元素': ws_tl.cell(row=row, column=5).value or '',
        '文案/字幕': ws_tl.cell(row=row, column=6).value or '',
        '音效': ws_tl.cell(row=row, column=7).value or '',
        '目的': ws_tl.cell(row=row, column=8).value or '',
        '可借鉴点': ws_tl.cell(row=row, column=9).value or '',
    })

if tl_rows:
    ok = 0
    for i in range(0, len(tl_rows), 20):
        batch = tl_rows[i:i+20]
        ok += insert_records(T_TIMELINE, batch)
        time.sleep(0.3)
    print(f'  逐帧拆解: {ok}/{len(tl_rows)} 条')

print('\n=== 同步总结分析 ===')
ws_sum = wb['总结分析']
sum_rows = []
for row in range(2, ws_sum.max_row + 1):
    name = ws_sum.cell(row=row, column=1).value
    if not name:
        continue
    record = {
        '素材名称': ws_sum.cell(row=row, column=1).value or '',
        '游戏名称': ws_sum.cell(row=row, column=2).value or '',
        '素材类型': ws_sum.cell(row=row, column=3).value or '',
        'Hook内容': ws_sum.cell(row=row, column=4).value or '',
        'Hook类型': ws_sum.cell(row=row, column=5).value or '',
        '人物': ws_sum.cell(row=row, column=6).value or '',
        'UI元素': ws_sum.cell(row=row, column=7).value or '',
        '特效': ws_sum.cell(row=row, column=8).value or '',
        '色彩': ws_sum.cell(row=row, column=9).value or '',
        '字幕': ws_sum.cell(row=row, column=10).value or '',
        '节奏': ws_sum.cell(row=row, column=11).value or '',
        'BGM': ws_sum.cell(row=row, column=12).value or '',
        '核心卖点': ws_sum.cell(row=row, column=13).value or '',
        'hook_策略': ws_sum.cell(row=row, column=14).value or '',
        '角色元素': ws_sum.cell(row=row, column=15).value or '',
        '场景元素': ws_sum.cell(row=row, column=16).value or '',
        '制作方式': ws_sum.cell(row=row, column=17).value or '',
        '叙事模式': ws_sum.cell(row=row, column=18).value or '',
        '画面风格': ws_sum.cell(row=row, column=19).value or '',
        '目标受众': ws_sum.cell(row=row, column=20).value or '',
        '情绪触发': ws_sum.cell(row=row, column=21).value or '',
        '营销策略': ws_sum.cell(row=row, column=22).value or '',
        '可优化方向': ws_sum.cell(row=row, column=23).value or '',
        '标签状态': ws_sum.cell(row=row, column=24).value or '',
        'AI置信度': ws_sum.cell(row=row, column=25).value or 0,
        '分析时间': ws_sum.cell(row=row, column=26).value or '',
    }
    sum_rows.append(record)

if sum_rows:
    ok = 0
    for i in range(0, len(sum_rows), 20):
        batch = sum_rows[i:i+20]
        ok += insert_records(T_ANALYSIS, batch)
        time.sleep(0.3)
    print(f'  AI分析结果: {ok}/{len(sum_rows)} 条')

wb.close()
print('\n✅ 同步完成')
