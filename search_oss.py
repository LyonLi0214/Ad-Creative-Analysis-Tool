import oss2

auth = oss2.Auth('[REDACTED]', '[REDACTED]')
bucket = oss2.Bucket(auth, 'https://oss-cn-shanghai.aliyuncs.com', 'mova-eu-game-uxd')

prefix = '心动防线-日/'
mp4_files = []
img_files = []
other_count = 0

for obj in oss2.ObjectIterator(bucket, prefix=prefix):
    key = obj.key
    if key.endswith('/'):
        continue
    lo = key.lower()
    if lo.endswith('.mp4'):
        mp4_files.append((key, obj.size))
    elif lo.endswith(('.png', '.jpg', '.jpeg', '.gif')):
        img_files.append((key, obj.size))
    else:
        other_count += 1

print(f'=== 心动防线-日/ ===')
print(f'MP4: {len(mp4_files)}')
print(f'图片: {len(img_files)}')
print(f'其他: {other_count}')

print(f'\n--- MP4 文件 ---')
for key, size in mp4_files[:30]:
    name = key.split('/')[-1]
    print(f'  {name}  ({size/1024/1024:.1f}MB)')
if len(mp4_files) > 30:
    print(f'  ... and {len(mp4_files)-30} more')

# 看看有没有和素材名匹配的
print(f'\n--- 匹配素材名 ---')
from openpyxl import load_workbook
ROOT = Path(__file__).resolve().parent
wb = load_workbook(ROOT / '素材拆解分析模板.xlsx', data_only=True)
ws = wb['素材总览']
asset_names = []
for row in range(2, ws.max_row + 1):
    name = ws.cell(row=row, column=1).value
    if name:
        asset_names.append(str(name).strip())

# 匹配逻辑：素材名的关键词在 OSS key 中
oss_names = {key.split('/')[-1]: key for key, _ in mp4_files}
matched = 0
for asset in asset_names[:10]:
    # 去掉 .mp4 后缀
    clean = asset.replace('.mp4', '')
    found = False
    for oss_name, oss_key in oss_names.items():
        if clean in oss_name or oss_name.replace('.mp4', '') in clean:
            print(f'  MATCH: {asset} -> {oss_key}')
            matched += 1
            found = True
            break
    if not found:
        print(f'  MISS:  {asset}')

print(f'\n前10条中匹配: {matched}/10')
