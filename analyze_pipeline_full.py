#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
analyze_pipeline_full.py — 完整分析管线（正式版）

流程: OSS 下载 → 本地识别(OCR/ASR/CLIP) → evidence JSON → AI 分析 → 写入 Excel+SQLite

用法:
    python analyze_pipeline_full.py --game 燃烧纪元              # 指定游戏
    python analyze_pipeline_full.py --game 心动防线              # 心动防线（日本）
    python analyze_pipeline_full.py --game 燃烧纪元 --limit 3    # 试跑3条
"""
import sys
sys.stdout.reconfigure(encoding='utf-8')

import os, json, re, time, subprocess, sqlite3, requests, hashlib, argparse
from pathlib import Path
from datetime import datetime, timezone
from openpyxl import load_workbook
from collections import defaultdict

ROOT = Path(__file__).parent
XLSX_PATH = ROOT / "素材拆解分析模板.xlsx"
DB_PATH = ROOT / "data" / "mova_materials.sqlite3"
DOWNLOAD_DIR = ROOT / "downloads" / "new_materials"
LABEL_SCHEMA_PATH = ROOT / "label_schema.yaml"
RECOGNITION_VENV = ROOT / ".venv-recognition" / "Scripts" / "python.exe"

# 游戏配置
GAME_CONFIG = {
    "心动防线": {"lang": "ja", "region": "日本"},
    "燃烧纪元": {"lang": "zh", "region": "港澳台"},
}

# ── 加载 .env ──
env_path = ROOT / ".env"
for line in env_path.read_text(encoding='utf-8-sig').splitlines():
    line = line.strip()
    if not line or line.startswith('#'): continue
    k, _, v = line.partition('=')
    os.environ.setdefault(k.strip(), v.strip())

OPENAI_BASE = os.environ.get("OPENAI_BASE", "https://api.deepseek.com/v1")
OPENAI_KEY = os.environ.get("OPENAI_KEY", "")
OPENAI_MODEL = os.environ.get("OPENAI_MODEL", "deepseek-v4-flash")

def material_key(name):
    return "mat_" + hashlib.md5(name.encode()).hexdigest()[:16]

def load_env():
    """确保 .env 已加载"""
    pass

# ═══════════════════════════════════════════════════════════
# 1. OSS 下载
# ═══════════════════════════════════════════════════════════
def download_from_oss(material_names: list[str], prefix: str) -> dict[str, Path]:
    """从 OSS 下载素材到 new_materials/，返回 {原始素材名: 本地路径}"""
    import oss2
    auth = oss2.Auth(os.environ['OSS_AK_ID'], os.environ['OSS_AK_SECRET'])
    bucket = oss2.Bucket(auth, os.environ['OSS_ENDPOINT'], os.environ['OSS_BUCKET'])

    print("构建 OSS 索引...")
    oss_index = {}
    for obj in oss2.ObjectIterator(bucket, prefix=prefix):
        if obj.key.endswith('/'): continue
        fname = obj.key.split('/')[-1]
        if fname not in oss_index: oss_index[fname] = []
        oss_index[fname].append(obj.key)

    def match(name: str) -> str | None:
        clean = name.strip().replace('.mp4', '')
        if clean + '.mp4' in oss_index: return oss_index[clean + '.mp4'][0]
        if clean in oss_index: return oss_index[clean][0]
        for fn, keys in oss_index.items():
            fnc = fn.replace('.mp4', '')
            if clean in fnc or fnc in clean: return keys[0]
        m_id = re.search(r'((?:V|P)\d+)', clean)
        m_date = re.search(r'(\d{8})', clean)
        m_res = re.search(r'(9x16|1x1|16x9)', clean)
        best, best_score = None, 0
        for fn, keys in oss_index.items():
            score = 0
            if m_id and m_id.group(1) in fn: score += 3
            if m_date and m_date.group(1) in fn: score += 2
            if m_res and m_res.group(1).lower() in fn.lower(): score += 1
            if score > best_score: best, best_score = keys[0], score
        return best if best_score >= 5 else None

    result = {}
    for mn in material_names:
        oss_key = match(mn)
        if not oss_key:
            print(f"  [MISS] {mn[:40]}")
            continue
        fname = oss_key.split('/')[-1]
        local = DOWNLOAD_DIR / fname
        if not local.exists():
            print(f"  下载 {fname}...", end=" ", flush=True)
            bucket.get_object_to_file(oss_key, str(local))
            print(f"[{local.stat().st_size/1024/1024:.1f}MB]")
        result[mn] = local
    return result

# ═══════════════════════════════════════════════════════════
# 2. 本地识别
# ═══════════════════════════════════════════════════════════
def run_local_recognition(video_path: Path, lang: str) -> Path | None:
    work_dir = video_path.parent / f"{video_path.stem}_work"
    evidence_path = video_path.parent / f"{video_path.stem}_evidence.json"
    if evidence_path.exists():
        print(f"  evidence 已存在")
        return evidence_path
    if not RECOGNITION_VENV.exists():
        print(f"  SKIP: 识别 venv 未找到")
        return None

    print(f"  本地识别 ({lang})...", end=" ", flush=True)
    t0 = time.time()
    try:
        result = subprocess.run(
            [str(RECOGNITION_VENV), "-m", "local_recognition.recognize_asset",
             str(video_path), "--work-dir", str(work_dir),
             "--output", str(evidence_path), "--language", lang],
            capture_output=True, text=True, timeout=600,
            encoding="utf-8", errors="replace",
        )
        elapsed = time.time() - t0
        if evidence_path.exists():
            print(f"[OK] {elapsed:.0f}s")
            return evidence_path
        else:
            print(f"[FAIL] {result.stderr[:150]}")
            return None
    except subprocess.TimeoutExpired:
        print("[TIMEOUT]")
        return None
    except Exception as e:
        print(f"[ERR] {e}")
        return None

# ═══════════════════════════════════════════════════════════
# 3. Evidence → Prompt
# ═══════════════════════════════════════════════════════════
def load_game_bg(game_name: str) -> str:
    try:
        db = sqlite3.connect(str(DB_PATH))
        row = db.execute(
            "SELECT product_background, core_gameplay, selling_points, theme_style, target_market FROM game_background WHERE game_name=?",
            (game_name,)
        ).fetchone()
        db.close()
        if row:
            return (f"产品背景: {row[0]}\n核心玩法: {row[1]}\n"
                    f"卖点: {row[2]}\n美术风格: {row[3]}\n目标市场: {row[4]}")
    except: pass
    return ""

def summarize_evidence(evidence_path: Path) -> str:
    with open(evidence_path, "r", encoding="utf-8") as f:
        ev = json.load(f)
    lines = []
    media = ev.get("media", {})
    lines.append(f"素材类型: {ev.get('asset_type', '?')}")
    lines.append(f"时长: {media.get('duration_sec', '?')}s | 分辨率: {media.get('width', '?')}x{media.get('height', '?')}")
    lines.append(f"有音频: {media.get('has_audio', False)} | 片段数: {len(ev.get('segments', []))}")
    ocr_texts = ev.get("global_evidence", {}).get("ocr_full_text", [])
    if ocr_texts:
        lines.append(f"\nOCR 识别文本 ({len(ocr_texts)} 条):")
        lines.append("; ".join(ocr_texts[:30]))
    speech = ev.get("global_evidence", {}).get("speech_full_text", [])
    if speech:
        lines.append(f"\nASR 语音转写 ({len(speech)} 条):")
        lines.append("; ".join(speech[:20]))
    cta = ev.get("global_evidence", {}).get("cta_present", False)
    lines.append(f"\nCTA 检测: {'有转化引导' if cta else '未检测到'}")
    lines.append(f"\n逐段拆解 ({len(ev.get('segments', []))} 段):")
    for seg in ev.get("segments", []):
        sid = seg.get("segment_id", "?")
        start = seg.get("start_sec", 0)
        end = seg.get("end_sec", 0)
        seg_ocr = [o["text_original"] for o in seg.get("ocr", [])[:5]]
        seg_speech = [s["text_original"] for s in seg.get("speech", [])[:3]]
        visual = seg.get("visual", {}).get("semantic_candidates", [])
        v_labels = []
        for v in visual:
            tp = "=" if v.get("threshold_passed") else "≈"
            v_labels.append(f"{v['group']}{tp}{v['top_label_zh']}({v['relative_score']:.0%})")
        lines.append(f"  {sid} [{start:.0f}-{end:.0f}s]:")
        if seg_ocr: lines.append(f"    OCR: {' | '.join(seg_ocr)}")
        if seg_speech: lines.append(f"    ASR: {' | '.join(seg_speech)}")
        if v_labels: lines.append(f"    CLIP: {', '.join(v_labels[:3])}")
    return "\n".join(lines)

def load_label_tags():
    import yaml
    with open(LABEL_SCHEMA_PATH, "r", encoding="utf-8") as f:
        schema = yaml.safe_load(f)
    dims = schema.get("dimensions", {})
    lines = ["## 六维标签体系"]
    for dn, di in dims.items():
        vs = "、".join(v["name"] for v in di.get("values", []))
        lines.append(f"- {dn}（{di.get('description', '')}）：{vs}")
    return "\n".join(lines)

def build_prompt(evidence_text, material_name, game_name, game_bg, tags_text, region):
    cfg = GAME_CONFIG.get(game_name, {})
    lang_hint = "繁体中文" if cfg.get("lang") == "zh" else "日文"

    system = f"""你是游戏广告素材分析专家。你收到的是经过本地 OCR/ASR/CLIP 视觉模型预处理后的素材证据摘要。请基于这些证据进行「四维拆解」并从六维标签体系中选择标签。

## 游戏背景
{game_bg or "（无）"}

## 四维拆解框架
1. **前3秒·钩子**：开头用什么方式抓住注意力？
2. **核心卖点·主体**：主体展示了什么核心内容？
3. **视听语言·包装**：画面风格、BGM、节奏如何？
4. **转化路径·CTA**：结尾如何引导转化？

{tags_text}

## 输出格式
只输出一个 JSON 对象（不要 markdown，不要多余文字）：
{{
  "逐帧拆解": [
    {{"时间轴":"0:00-0:03","画面内容":"...","画面元素":"...","文案/字幕":"OCR/ASR原文放这里","音效":"...","目的":"...","可借鉴点":"..."}}
  ],
  "总结": {{
    "Hook内容":"...","Hook类型":"...","人物":"...","UI元素":"...","特效":"...","色彩":"...","字幕":"...","节奏":"...","BGM":"...","核心卖点":"...",
    "hook_策略":"...","角色元素":"...","场景元素":"...","制作方式":"...","叙事模式":"...","画面风格":"...",
    "目标受众":"...","情绪触发":"...","营销策略":"...","可优化方向":"...","AI置信度":0.8
  }}
}}

## 规则
- 标签值必须从六维标签体系中选取，不得自创
- **文案/字幕 字段必须放入该段真实 OCR/ASR 原文**，这是给 copywriter 做变体用的
- 画面文字语言为 {lang_hint}，OCR 结果如有乱码尽量还原
- CLIP 标签中带 ≈ 号的表示低于置信度阈值，仅供参考
- 证据中未出现的元素填"未识别到"
- AI置信度按证据充分程度给 0.0-1.0"""

    user = f"""素材名称: {material_name}
游戏: {game_name}（{region}）

## 本地识别证据
{evidence_text}

请基于以上证据进行分析。"""
    return system, user

# ═══════════════════════════════════════════════════════════
# 4. AI 分析
# ═══════════════════════════════════════════════════════════
def call_ai(system: str, user: str) -> dict | None:
    url = f"{OPENAI_BASE}/chat/completions"
    headers = {"Authorization": f"Bearer {OPENAI_KEY}", "Content-Type": "application/json"}
    body = {"model": OPENAI_MODEL, "messages": [
        {"role": "system", "content": system},
        {"role": "user", "content": user},
    ], "max_tokens": 8000, "temperature": 0.1}
    try:
        r = requests.post(url, headers=headers, json=body, timeout=300)
        r.raise_for_status()
        content = r.json()["choices"][0]["message"]["content"]
        if "```json" in content: content = content.split("```json")[1].split("```")[0]
        elif "```" in content and content.count("```") >= 2:
            content = content.split("```")[1].split("```")[0]
        content = re.sub(r',\s*}', '}', content.strip())
        content = re.sub(r',\s*]', ']', content)
        return json.loads(content)
    except json.JSONDecodeError:
        try:
            start = content.find('{'); end = content.rfind('}')
            if start >= 0 and end > start:
                s = content[start:end+1]
                s = re.sub(r',\s*}','}',s); s = re.sub(r',\s*]',']',s)
                return json.loads(s)
        except: pass
        print(f"    JSON 解析失败: {content[:200]}")
        return None
    except Exception as e:
        print(f"    API 失败: {e}")
        return None

# ═══════════════════════════════════════════════════════════
# 5. 写入 Excel + SQLite
# ═══════════════════════════════════════════════════════════
def write_to_excel(wb, name, game, analysis):
    summary = analysis.get("总结", analysis)
    timeline = analysis.get("逐帧拆解", [])
    ws = wb["总结分析"]
    hd = {ws.cell(row=1,column=c).value: c for c in range(1, ws.max_column+1)}
    rn = ws.max_row + 1
    ws.cell(row=rn, column=hd.get("素材名称",1)).value = name
    ws.cell(row=rn, column=hd.get("游戏名称",2)).value = game
    field_map = {"Hook内容":"Hook内容","Hook类型":"Hook类型","人物":"人物","UI元素":"UI元素",
        "特效":"特效","色彩":"色彩","字幕":"字幕","节奏":"节奏","BGM":"BGM","核心卖点":"核心卖点",
        "hook_策略":"hook_策略","角色元素":"角色元素","场景元素":"场景元素","制作方式":"制作方式",
        "叙事模式":"叙事模式","画面风格":"画面风格","目标受众":"目标受众","情绪触发":"情绪触发",
        "营销策略":"营销策略","可优化方向":"可优化方向","AI置信度":"AI置信度"}
    for eng, chi in field_map.items():
        v = summary.get(eng) or summary.get(chi, "")
        c = hd.get(chi)
        if c: ws.cell(row=rn, column=c).value = str(v) if v else ""

    ws_tl = wb["逐帧拆解"]
    tl_hd = {ws_tl.cell(row=1,column=c).value: c for c in range(1, ws_tl.max_column+1)}
    for seg in timeline:
        r = ws_tl.max_row + 1
        ws_tl.cell(row=r, column=tl_hd.get("素材名称",1)).value = name
        ws_tl.cell(row=r, column=tl_hd.get("游戏名称",2)).value = game
        ws_tl.cell(row=r, column=tl_hd.get("时间轴",3)).value = seg.get("时间轴","")
        ws_tl.cell(row=r, column=tl_hd.get("画面内容",4)).value = seg.get("画面内容","")
        ws_tl.cell(row=r, column=tl_hd.get("画面元素",5)).value = seg.get("画面元素","")
        ws_tl.cell(row=r, column=tl_hd.get("文案/字幕",6)).value = seg.get("文案/字幕","")
        ws_tl.cell(row=r, column=tl_hd.get("音效",7)).value = seg.get("音效","")
        ws_tl.cell(row=r, column=tl_hd.get("目的",8)).value = seg.get("目的","")
        ws_tl.cell(row=r, column=tl_hd.get("可借鉴点",9)).value = seg.get("可借鉴点","")

def save_sqlite(mk, name, game, analysis):
    conn = sqlite3.connect(str(DB_PATH))
    now = datetime.now(timezone.utc).isoformat()
    summary = analysis.get("总结", analysis)
    conn.execute("""INSERT OR REPLACE INTO ai_analysis
        (material_key,material_name,game_name,hook_content,hook_type,characters,
         ui_elements,effects,colors,subtitles,rhythm,bgm,core_selling_point,
         hook_strategy,character_elements,scene_elements,production_method,
         narrative_mode,visual_style,target_audience,emotional_trigger,
         marketing_strategy,optimization_direction,ai_confidence,analysis_time,updated_at)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (mk,name,game,
         str(summary.get("Hook内容","")),str(summary.get("Hook类型","")),str(summary.get("人物","")),
         str(summary.get("UI元素","")),str(summary.get("特效","")),str(summary.get("色彩","")),
         str(summary.get("字幕","")),str(summary.get("节奏","")),str(summary.get("BGM","")),
         str(summary.get("核心卖点","")),
         str(summary.get("hook_策略","")),str(summary.get("角色元素","")),str(summary.get("场景元素","")),
         str(summary.get("制作方式","")),str(summary.get("叙事模式","")),str(summary.get("画面风格","")),
         str(summary.get("目标受众","")),str(summary.get("情绪触发","")),
         str(summary.get("营销策略","")),str(summary.get("可优化方向","")),
         float(summary.get("AI置信度",0.5)), now, now))

    conn.execute("DELETE FROM frame_timeline WHERE material_key=?", (mk,))
    for seg in analysis.get("逐帧拆解", []):
        conn.execute("""INSERT INTO frame_timeline
            (material_key,material_name,game_name,timeline,frame_content,frame_elements,
             subtitles,sound_effect,purpose,learnable_point,updated_at)
            VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
            (mk,name,game,seg.get("时间轴",""),seg.get("画面内容",""),
             seg.get("画面元素",""),seg.get("文案/字幕",""),seg.get("音效",""),
             seg.get("目的",""),seg.get("可借鉴点",""), now))
    conn.commit(); conn.close()

# ═══════════════════════════════════════════════════════════
# 6. 分析单条素材
# ═══════════════════════════════════════════════════════════
def analyze_one(material_name: str, game_name: str, local_file: Path) -> bool:
    mk = material_key(material_name)
    cfg = GAME_CONFIG.get(game_name, {"lang": "ja", "region": game_name})
    lang = cfg["lang"]
    region = cfg["region"]

    print(f"\n{'='*50}")
    print(f"分析: {material_name[:50]}")
    print(f"{'='*50}")

    # 1. 本地识别
    evidence_path = run_local_recognition(local_file, lang)
    if not evidence_path:
        print(f"  SKIP: 本地识别失败")
        return False

    # 2. Evidence 摘要
    evidence_text = summarize_evidence(evidence_path)
    print(f"  evidence: {len(evidence_text)} 字符")

    # 3. 游戏背景 + 标签
    game_bg = load_game_bg(game_name)
    tags_text = load_label_tags()

    # 4. AI 分析
    print(f"  AI ({OPENAI_MODEL})...", end=" ", flush=True)
    system, user = build_prompt(evidence_text, material_name, game_name, game_bg, tags_text, region)
    t0 = time.time()
    analysis = call_ai(system, user)
    if not analysis: print(f"[FAIL]"); return False
    print(f"[OK] {time.time()-t0:.0f}s")

    # 5. 写入 Excel
    print(f"  写入...", end=" ", flush=True)
    wb = load_workbook(str(XLSX_PATH))
    write_to_excel(wb, material_name, game_name, analysis)
    wb.save(str(XLSX_PATH)); wb.close()

    # 6. 保存 SQLite
    save_sqlite(mk, material_name, game_name, analysis)
    print("[OK]")
    return True

# ═══════════════════════════════════════════════════════════
# 主流程
# ═══════════════════════════════════════════════════════════
def main():
    parser = argparse.ArgumentParser(description="完整分析管线：OSS→本地识别→AI→Excel")
    parser.add_argument("--game", required=True, help="游戏名称")
    parser.add_argument("--limit", type=int, default=999, help="最多分析条数")
    parser.add_argument("--oss-prefix", help="OSS 前缀（默认根据游戏名自动匹配）")
    args = parser.parse_args()

    game = args.game
    cfg = GAME_CONFIG.get(game, {"lang": "ja", "region": game})

    # 自动 OSS 前缀
    oss_prefix = args.oss_prefix
    if not oss_prefix:
        PREFIX_MAP = {"心动防线": "心动防线-日/", "燃烧纪元": "燃烧纪元-投放素材/"}
        oss_prefix = PREFIX_MAP.get(game, f"{game}/")

    DOWNLOAD_DIR.mkdir(parents=True, exist_ok=True)

    # 从 Excel 读待分析素材
    wb = load_workbook(str(XLSX_PATH))
    ws = wb["素材总览"]
    col_map = {}
    for c in range(1, ws.max_column + 1):
        h = ws.cell(row=1, column=c).value
        if h: col_map[str(h).strip()] = c

    materials = []
    for r in range(2, ws.max_row + 1):
        g = ws.cell(row=r, column=col_map.get("游戏名称", 2)).value
        s = ws.cell(row=r, column=col_map.get("状态", 9)).value
        if g and game in str(g) and (not s or str(s).strip() == "待分析"):
            name = str(ws.cell(row=r, column=col_map.get("素材名称", 1)).value or "").strip()
            if name: materials.append(name)
    wb.close()

    # 去重（去掉 .mp4 后缀比较）
    seen = set()
    unique = []
    for m in materials:
        key = m.replace('.mp4', '').strip()
        if key not in seen: seen.add(key); unique.append(m)
    materials = unique[:args.limit]

    print(f"游戏: {game} ({cfg.get('region', '?')})")
    print(f"语言: {cfg.get('lang', 'ja')}")
    print(f"OSS: {oss_prefix}")
    print(f"AI: {OPENAI_MODEL}")
    print(f"待分析: {len(materials)} 条")
    print()

    if not materials:
        print("没有待分析素材。")
        return

    # 下载
    print("=== OSS 下载 ===")
    local_files = download_from_oss(materials, oss_prefix)

    # 分析
    print("\n=== 分析 ===")
    results = {"success": 0, "fail": 0}
    for i, mn in enumerate(materials, 1):
        if mn not in local_files:
            print(f"\n[{i}/{len(materials)}] {mn[:40]} [MISS 无下载文件]")
            results["fail"] += 1
            continue
        print(f"\n[{i}/{len(materials)}]")
        ok = analyze_one(mn, game, local_files[mn])
        if ok: results["success"] += 1
        else: results["fail"] += 1

    print(f"\n{'='*50}")
    print(f"完成: {results['success']} 成功, {results['fail']} 失败")
    print(f"提示: 导出专用Excel → 运行 export_game_excel.py --game {game}")
    print(f"{'='*50}")

if __name__ == "__main__":
    main()
